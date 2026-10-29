#!/usr/bin/env python3
"""
End-to-end test: every collector against the mock vendor APIs, into a
throwaway Postgres database, then the rollups, then every SQL query in the
Grafana dashboards. Exits non-zero on any failure.

  PGHOST=127.0.0.1 PGUSER=postgres python3 tests/run_e2e.py

Needs a Postgres it may DROP/CREATE the database `secops_e2e` on.
"""
import datetime as dt
import glob
import json
import os
import re
import subprocess
import sys
import tempfile

import psycopg2

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
sys.path.insert(0, HERE)

import mock_vendors  # noqa: E402

PG = dict(host=os.environ.get("PGHOST", "127.0.0.1"), port=int(os.environ.get("PGPORT", 5432)),
          user=os.environ.get("PGUSER", "postgres"), password=os.environ.get("PGPASSWORD"))
DB = "secops_e2e"
PORT = 8999
M = f"http://127.0.0.1:{PORT}"

CONFIG = f"""
crowdstrike: {{base_url: "{M}/falcon"}}
reporting: {{days_last_seen: 30, require_exploit_for_remote_no_auth: true, stale_sensor_days: 7,
            sla_days: {{critical: 2, high: 14, medium: 30, low: 60}}}}
sites:
  - key: RVH
    label: Riverside Health
    match: {{falcon_groups: ["RVH Workstations"], ad_domains: ["rvh.local"], email_domains: ["riversidehealth.test"],
            cidrs: ["10.1.0.0/16", "203.0.113.0/28"]}}
  - key: LKH
    label: Lakeside Health
    match: {{falcon_tags: ["SensorGroupingTags/LKH"], ou_contains: ["ou=lakeside"],
            email_domains: ["lakesidehealth.test"], hostname_regex: ["^LKH-"]}}
ungrouped_label: Ungrouped
database: {{host: "{PG['host']}", port: {PG['port']}, name: {DB}, user: "{PG['user']}"}}
secrets_file: secrets.yaml
collectors:
  falcon_hosts: {{enabled: true, out_of_scope_ous: {{"Departed Service": "left the alliance"}}}}
  falcon_spotlight: {{enabled: true, page_size: 2}}
  falcon_alerts: {{enabled: true}}
  falcon_identity: {{enabled: true, access_domain: {{domain: "shared.local", native_sites: ["Lakeside Health"]}}}}
  falcon_usb: {{enabled: true, lookback_days: 5}}
  hadrian:
    enabled: true
    base_url: "{M}/hadrian"
    organization_id: "org1"
    tag_aliases: {{OLDRVH: RVH}}
    archive_tags: ["zzArchive"]
    suggested_tags_out: "hadrian-suggested-tags.csv"
    detail_workers: 2
    assets: {{path: "/organizations/{{organization_id}}/assets", items_key: items,
             pagination: {{type: page, param: offset, size_param: pageSize, size: 1, start: 0}},
             fields: {{id: assetId, name: value, type: platformAssetType,
                      first_seen: detectedOnUtc, last_seen: lastSeenAtUtc}}}}
    risks: {{path: "/organizations/{{organization_id}}/risks", detail_path: "/organizations/{{organization_id}}/risks/{{id}}",
            items_key: items, pagination: {{type: page, param: offset, size_param: pageSize, size: 1, start: 0}},
            fields: {{id: id, title: title, severity: riskSeverity, activity: activityStatus, status: status,
                     visibility: riskVisibility, risk_type: riskType, category: primaryCategory.id,
                     first_seen: created, last_seen: lastSeen, resolved_at: resolvedOn}}}}
  checkpoint_hec: {{enabled: true, gateway: "{M}/hec"}}
  checkpoint_hec_flow: {{enabled: true, lookback_days: 1}}
  entra:
    enabled: true
    tenants: [{{name: main, tenant_id: t1, login_base: "{M}/login", graph_base: "{M}/graph",
               loganalytics_base: "{M}/la"}}]
  azure_log_analytics:
    enabled: true
    workspaces: [{{name: ws, tenant: main, workspace_id: w1}}]
    queries: [{{name: signins, key_type: upn, kql: "SigninLogs | take 1"}}]
  dmarc: {{enabled: true, opensearch_url: "{M}/os", username: u}}
notify: {{enabled: true, host_label: e2e, healthcheck_url: "{M}/hc/abc",
          pushover: {{api_url: "{M}/pushover/1/messages.json"}}}}
"""
SECRETS = """
database: {password: null}
crowdstrike: {client_id: cid, client_secret: sec}
collectors:
  hadrian: {api_key: hk}
  checkpoint_hec: {client_id: hc, access_key: hk}
  entra: {tenants: [{name: main, client_id: ac, client_secret: as}]}
  dmarc: {password: p}
notify: {pushover: {app_token: at, user_key: uk}}
"""

failures = []


def check(cond, msg):
    print(("  PASS " if cond else "  FAIL ") + msg)
    if not cond:
        failures.append(msg)


def q(cur, sql, *a):
    cur.execute(sql, a)
    return cur.fetchall()


def main() -> int:
    admin = psycopg2.connect(dbname="postgres", **PG)
    admin.autocommit = True
    admin.cursor().execute(f"DROP DATABASE IF EXISTS {DB}")
    admin.cursor().execute(f"CREATE DATABASE {DB}")

    tmp = tempfile.mkdtemp()
    open(os.path.join(tmp, "config.yaml"), "w").write(CONFIG)
    open(os.path.join(tmp, "secrets.yaml"), "w").write(SECRETS)
    cfgp = os.path.join(tmp, "config.yaml")
    srv = mock_vendors.start(PORT)

    def run(*args):
        r = subprocess.run([sys.executable, *args, "--config", cfgp], cwd=ROOT, capture_output=True, text=True)
        print(r.stdout[-4000:])
        if r.returncode != 0:
            print(r.stderr[-4000:])
        return r.returncode

    # schema first, then seed KEV/EPSS and a stale open row that must get EXPIRED
    import config as config_mod
    import db_schema
    cfg = config_mod.load_config(cfgp)
    db_schema.ensure_schema(cfg)
    conn = psycopg2.connect(dbname=DB, **PG)
    cur = conn.cursor()
    cur.execute("INSERT INTO cisa_kev (cve_id, date_added, due_date, known_ransomware_campaign_use) "
                "VALUES ('CVE-2021-44228', '2021-12-10', '2021-12-24', 'Known')")
    cur.execute("INSERT INTO epss_scores (cve_id, epss, percentile) VALUES ('CVE-2024-0001', 0.71, 0.99)")
    cur.execute("""INSERT INTO vuln_findings (source, source_asset_id, source_rule_id, state, severity, first_found,
                   last_found, site_label, last_seen_in_export)
                   VALUES ('falcon_spotlight','aid9','gone','OPEN','high', now()-interval '50 days',
                           now()-interval '1 day','Ungrouped', now()-interval '1 day')""")
    conn.commit()

    print("== collect.py (all collectors)")
    rc = run("collect.py")
    check(rc == 0, "collect.py exited 0")
    print("== collect.py second run (idempotency / incremental)")
    check(run("collect.py") == 0, "second collect.py run exited 0")
    # incremental open pull: only changed vulns asked for, and nothing expired (the pull isn't complete)
    cfg_inc = os.path.join(tmp, "config-incremental.yaml")
    open(cfg_inc, "w").write(CONFIG.replace("falcon_spotlight: {enabled: true, page_size: 2}",
                                            "falcon_spotlight: {enabled: true, page_size: 2, open_pull: incremental, "
                                            "full_pull_weekdays: []}"))
    cur.execute("""INSERT INTO vuln_findings (source, source_asset_id, source_rule_id, state, severity, first_found,
                   last_found, site_label, last_seen_in_export)
                   VALUES ('falcon_spotlight','aid9','gone2','OPEN','high', now()-interval '50 days',
                           now()-interval '1 day','Ungrouped', now()-interval '1 day')""")
    conn.commit()
    n_calls = len(mock_vendors.CALLS)
    r = subprocess.run([sys.executable, "collect.py", "--config", cfg_inc, "--only", "falcon_spotlight"],
                       cwd=ROOT, capture_output=True, text=True)
    check(r.returncode == 0, "incremental Spotlight run exited 0")
    if r.returncode != 0:
        print(r.stdout[-2000:], r.stderr[-2000:])
    flts = [c[2].get("filter", [""])[0] for c in mock_vendors.CALLS[n_calls:]
            if c[1] == "/falcon/spotlight/combined/vulnerabilities/v1" and "after" not in c[2]]
    check(len(flts) == 2 and all("updated_timestamp:>" in f for f in flts),
          f"incremental run asks only for updated open + closed vulns: {flts}")
    mode = q(cur, "SELECT stats->>'open_pull' FROM collector_runs WHERE collector = 'falcon_spotlight' "
                  "ORDER BY started_at DESC LIMIT 1")[0][0]
    st2 = q(cur, "SELECT state FROM vuln_findings WHERE source_rule_id = 'gone2'")[0][0]
    check(mode == "incremental" and st2 == "OPEN", f"incremental pull expires nothing ({mode}, gone2 {st2})")
    cur.execute("DELETE FROM vuln_findings WHERE source_rule_id = 'gone2'")
    conn.commit()
    print("== rollup_daily_metrics.py")
    check(run("rollup_daily_metrics.py") == 0, "rollup exited 0")

    print("== notifications")
    # form-encoded, so the mock parses each field to a list
    pushes = lambda: [{k: v[0] for k, v in c[2].items()} for c in mock_vendors.CALLS
                      if c[1] == "/pushover/1/messages.json"]
    pings = lambda: [c[1] for c in mock_vendors.CALLS if c[1].startswith("/hc/")]
    started = (dt.datetime.now().astimezone() - dt.timedelta(minutes=5)).isoformat(timespec="seconds")
    subprocess.run([sys.executable, "notify.py", "--config", cfgp, "run-start"], cwd=ROOT)
    subprocess.run([sys.executable, "notify.py", "--config", cfgp, "run-end", "--rc", "0", "--started", started], cwd=ROOT)
    check(pings() == ["/hc/abc/start", "/hc/abc/0"] and not pushes(),
          f"healthy run: start + success pings, no Pushover message ({pings()}, {len(pushes())} pushes)")
    cur.execute("INSERT INTO collector_runs (collector, started_at, finished_at, status, error) "
                "VALUES ('falcon_alerts', now(), now(), 'error', 'HTTP 403 from vendor')")
    conn.commit()
    subprocess.run([sys.executable, "notify.py", "--config", cfgp, "run-end", "--rc", "1", "--started", started], cwd=ROOT)
    p = pushes()
    check(len(p) == 1 and "falcon_alerts: error -- HTTP 403" in p[0].get("message", "")
          and p[0].get("priority") == "1" and pings()[-1] == "/hc/abc/1",
          f"failed collector: one high-priority push naming it, failure ping ({p}, {pings()[-1:]})")
    cur.execute("DELETE FROM collector_runs WHERE error = 'HTTP 403 from vendor'")
    # source gone quiet: the collector still succeeds, but its newest data is a week old
    cur.execute("UPDATE dmarc_daily SET report_date = current_date - 7")
    conn.commit()
    subprocess.run([sys.executable, "notify.py", "--config", cfgp, "run-end", "--rc", "0", "--started", started], cwd=ROOT)
    p = pushes()
    check(len(p) == 2 and "dmarc: runs OK but newest data is" in p[1].get("message", "")
          and p[1].get("priority") == "0" and pings()[-1] == "/hc/abc/0",
          f"quiet feed: normal-priority push naming it, success ping ({p[1:]}, {pings()[-1:]})")
    srv.shutdown()

    cur = conn.cursor()
    print("== collector_runs")
    rows = q(cur, "SELECT collector, status, error FROM collector_freshness ORDER BY 1")
    for r in rows:
        check(r[1] == "ok", f"collector {r[0]} status ok ({r[2] or ''})")
    check(len(rows) == 11, f"11 collectors ran (got {len(rows)})")

    print("== site resolution")
    sites = dict(q(cur, "SELECT source_asset_id, site_label || '/' || site_matched_by FROM assets"))
    check(sites.get("aid1") == "Riverside Health/falcon_groups", f"aid1 via host group: {sites.get('aid1')}")
    check(sites.get("aid2") == "Lakeside Health/falcon_tags", f"aid2 via tag: {sites.get('aid2')}")
    check(sites.get("aid3") == "Riverside Health/ad_domains", f"aid3 via AD domain: {sites.get('aid3')}")
    check(sites.get("aid4") == "Ungrouped/none", f"aid4 ungrouped: {sites.get('aid4')}")
    check(sites.get("u1") == "Riverside Health/cidrs", f"unmanaged u1 via CIDR: {sites.get('u1')}")
    check(sites.get("u5") == "Lakeside Health/discoverer", f"network-only u5 via its discoverer's site: {sites.get('u5')}")
    cls = dict(q(cur, "SELECT source_asset_id, discovery_class FROM assets WHERE source = 'falcon_unmanaged'"))
    check(cls == {"u1": "network_device", "u2": "workstation_no_sensor", "u3": "service_account",
                  "u4": "duplicate_of_managed", "u5": "vmware_nic", "u6": "out_of_scope",
                  "u7": "secondary_ip_of_managed", "u8": "duplicate_of_managed"}, f"Discover classes {cls}")
    twin = q(cur, "SELECT managed_twin FROM assets WHERE source_asset_id = 'u4'")[0][0]
    check(twin == "RVH-WS01", f"duplicate linked to its managed host: {twin}")

    print("== vulnerabilities")
    st = dict(q(cur, "SELECT source_rule_id, state FROM vuln_findings"))
    check(st.get("v1") == "OPEN" and st.get("v2") == "REOPENED", "open/reopen mapped")
    check(st.get("v3") == "SUPPRESSED", "suppressed vuln -> SUPPRESSED")
    check(st.get("v5") == "FIXED", "closed vuln -> FIXED")
    check(st.get("gone") == "EXPIRED", "vanished open vuln -> EXPIRED (not FIXED)")
    check((st.get("r1"), st.get("r2"), st.get("r3")) == ("OPEN", "FIXED", "REOPENED"),
          f"Hadrian activity/status -> OPEN/FIXED(NotFound)/REOPENED: {st.get('r1'), st.get('r2'), st.get('r3')}")
    check("r4" not in st, "risk only on a zzArchive asset skipped")
    rt = dict(q(cur, "SELECT source_rule_id, risk_type FROM vuln_findings WHERE source = 'hadrian'"))
    check(rt.get("r1") == "UnpatchedTechnology" and rt.get("r3") == "InfectedDevice", f"risk_type stored: {rt}")
    xa = dict(q(cur, "SELECT asset_id, site_label || '/' || site_matched_by FROM external_assets"))
    check(xa == {"a1": "Riverside Health/hadrian_tag", "a2": "Lakeside Health/hadrian_apex"},
          f"external asset sites (tag, apex; archived skipped): {xa}")
    v1 = q(cur, "SELECT is_remote_no_auth, exploit_available, product_key, fix_id, site_label FROM vuln_findings WHERE source_rule_id='v1'")[0]
    check(v1 == (True, True, "google:chrome", "R1", "Riverside Health"), f"v1 enrichment {v1}")
    v2r = q(cur, "SELECT is_remote_no_auth FROM vuln_findings WHERE source_rule_id='v2'")[0][0]
    check(v2r is False, "no vector -> not remote/no-auth (old cvss bug fixed)")
    lf = q(cur, "SELECT last_found > now() - interval '1 day' FROM vuln_findings WHERE source_rule_id='v1'")[0][0]
    check(lf, "last_found taken from host last_seen")
    runs = q(cur, "SELECT stats->>'rows', stats->>'changed' FROM collector_runs WHERE collector = 'falcon_spotlight' "
                  "AND stats->>'open_pull' = 'full' ORDER BY started_at")
    check(len(runs) == 2 and runs[1][1] == "0" and runs[1][0] != "0",
          f"second identical pull rewrites no findings (rows, changed per run: {runs})")
    desc = q(cur, "SELECT count(*) FROM cve_descriptions")[0][0]
    syn = q(cur, "SELECT count(*) FROM vuln_findings WHERE source = 'falcon_spotlight' AND synopsis IS NOT NULL")[0][0]
    check(desc >= 4 and syn == 0, f"CVE descriptions stored once per CVE ({desc}), not per finding ({syn})")
    kev = q(cur, "SELECT count(*) FROM fact_vuln_findings_current WHERE has_kev")[0][0]
    check(kev == 1, f"KEV joins Spotlight log4j; Hadrian risks carry no CVEs (got {kev})")
    patch = q(cur, "SELECT title, open_findings, affected_assets FROM patch_impact_summary WHERE fix_key='R1'")
    check(patch and patch[0][1] == 2, f"R1 patch groups both Chrome CVEs: {patch}")
    ext = q(cur, "SELECT site_label, asset_type FROM vuln_findings WHERE source='hadrian' AND source_rule_id='r1'")[0]
    check(ext == ("Riverside Health", "internet"), f"Hadrian risk site/type {ext}")

    print("== other domains")
    check(q(cur, "SELECT count(*) FROM security_alerts")[0][0] == 3, "3 alerts")
    al2 = q(cur, "SELECT site_label, closed_at IS NOT NULL FROM security_alerts WHERE alert_id='al2'")[0]
    check(al2 == ("Lakeside Health", True), f"user-only alert -> site via UPN, closed_at set: {al2}")
    al2d = q(cur, "SELECT status, disposition FROM security_alerts WHERE alert_id='al2'")[0]
    check(al2d == ("closed", "false_positive"),
          f"disposition preserved under collapsed status: {al2d}")
    ids = dict(q(cur, "SELECT entity_id, site_label FROM identity_entities"))
    check(ids == {"e1": "Riverside Health", "e2": "Lakeside Health", "e3": "Ungrouped", "e4": "Riverside Health",
                  "e5": "Lakeside Health"}, f"identity sites {ids}")
    en = dict(q(cur, "SELECT entity_id, (enabled, last_activity IS NOT NULL) FROM identity_entities "
                     "WHERE entity_id IN ('e3', 'e5')"))
    check(en == {"e3": "(,f)", "e5": "(f,t)"}, f"cloud-only account takes enabled/activity from Entra: {en}")
    ros = {r[0]: r[1:] for r in q(cur, "SELECT entity_id, title, department, account_expires IS NOT NULL, "
                                       "password_never_expires FROM identity_entities WHERE entity_id IN ('e1','e2')")}
    check(ros == {"e1": (None, None, False, True), "e2": ("Registered Nurse", "Ward 3", True, False)},
          f"roster fields: title, department, expiry, password-never-expires flag: {ros}")
    acc = dict(q(cur, "SELECT entity_id, access_account FROM identity_entities WHERE access_account IS NOT NULL"))
    check(acc == {"e4": "vpn"}, f"shared-domain VPN-only account tagged: {acc}")
    kinds = dict(q(cur, "SELECT entity_id, account_kind FROM identity_entities"))
    check(kinds.get("e1") == "service" and kinds.get("e2") == "human", f"account kinds (svc_ name -> service): {kinds}")
    check(q(cur, "SELECT count(*) FROM identity_risk_factors")[0][0] == 3, "3 identity risk factors (not doubled on rerun)")
    hec = dict(q(cur, "SELECT event_id, site_label FROM email_events"))
    check(hec == {"h1": "Riverside Health", "h2": "Lakeside Health", "h3": "Ungrouped"}, f"HEC recipient sites {hec}")
    dirs = dict(q(cur, "SELECT event_id, direction FROM email_events"))
    check(dirs == {"h1": "inbound", "h2": "inbound", "h3": "outbound"}, f"HEC direction classified: {dirs}")
    check(q(cur, "SELECT count(*) FROM entra_risky_users")[0][0] == 2, "2 risky users (paged via nextLink)")
    check(q(cur, "SELECT count(*) FROM entra_mfa_registration")[0][0] == 2, "guests excluded from MFA stats")
    la = dict(q(cur, "SELECT site_label, value FROM azure_log_metrics WHERE dimension='failure'"))
    check(la.get("Riverside Health") == 15, f"Log Analytics values summed per site: {la}")
    dm = q(cur, "SELECT sum(messages), sum(dmarc_pass), count(DISTINCT site_label) FROM dmarc_daily")[0]
    check(tuple(map(int, dm)) == (1537, 1500, 2), f"DMARC aggregates across scroll pages: {dm}")

    print("== removable media")
    usb = {(r[0], r[1]): r[2:] for r in q(cur, """SELECT user_key, category, files, bytes, site_label, folder,
                                                          display_name, device_serial FROM usb_file_writes_daily""")}
    check(sum(v[0] for v in usb.values()) == 6, f"6 USB file writes kept, Office lock file dropped: {usb}")
    check(usb.get(("jnurse", "email"), (0,))[0] == 1 and usb.get(("jnurse", "data"), (0,))[0] == 1,
          ".pst -> email, .xlsx -> data (extension beats the sensor's Generic label)")
    icons = usb.get(("jnurse", "images"))
    check(icons and icons[0] == 2 and icons[3] == "App\\Portable", f"program copy rolled up to one folder row: {icons}")
    check(usb.get(("jnurse", "data"), (0,) * 6)[2:4] == ("Riverside Health", "Ward\\Exports")
          and usb[("jnurse", "data")][4] == "Jo Nurse", "site by aid, folder, name from the enabled lookup row")
    nouser = usb.get(("", "archives"))
    check(nouser and nouser[2] == "Lakeside Health" and nouser[5] == "", f"no user / unknown drive kept: {nouser}")
    runs = q(cur, "SELECT stats->>'write_rows' FROM collector_runs WHERE collector = 'falcon_usb' ORDER BY started_at")
    check(len(runs) == 2 and runs[0] == runs[1], f"re-pull replaces days rather than doubling: {runs}")
    sc = dict(q(cur, "SELECT scan_id, site_label FROM usb_scans"))
    check(sc == {"sc_usb1": "Riverside Health", "sc_usb2": "Lakeside Health"}, f"USB-insertion scans by host site: {sc}")
    det = q(cur, """SELECT a.file_name, a.quarantined FROM security_alerts a
                    JOIN usb_scans u ON u.scan_id = a.scan_id""")
    check(det == [("setup.exe", True)], f"ODS detection linked to its USB scan: {det}")

    print("== mail-flow funnel")
    fl = {(r[0], r[1]): int(r[2]) for r in q(cur, "SELECT site_label, metric, sum(value) FROM daily_email_flow_metrics "
                                                 "GROUP BY 1, 2")}
    check((fl.get(("Riverside Health", "incoming")), fl.get(("Lakeside Health", "incoming")),
           fl.get(("Ungrouped", "incoming"))) == (4, 1, 1), f"incoming per site by recipient domain, rest Ungrouped: {fl}")
    check(fl.get(("Lakeside Health", "internal_incoming")) == 1 and fl.get(("Lakeside Health", "incoming")) == 1,
          "internal mail counted separately from incoming")
    check((fl.get(("Riverside Health", "ms_junk")), fl.get(("Riverside Health", "ms_quarantine")),
           fl.get(("Riverside Health", "cp_quarantined")), fl.get(("Riverside Health", "restored"))) == (1, 1, 3, 1),
          "Microsoft junk/quarantine, Check Point quarantine and restores counted")

    print("== rollups")
    am = dict(q(cur, "SELECT site_label, stale_sensors FROM daily_asset_metrics WHERE snapshot_date=CURRENT_DATE"))
    check(am.get("Riverside Health") == 1, f"stale sensor counted (aid3): {am}")
    sil = {r[0]: r[1:] for r in q(cur, "SELECT site_label, silent_sensors, silent_in_use FROM daily_asset_metrics "
                                       "WHERE snapshot_date = CURRENT_DATE")}
    # aid3: workstation silent 20 days -- inside the 30-day limit, but still logging on (u8), so not reporting;
    # aid2 (server, 2 days) and aid4 (DC, 1 day) are inside their limits
    check(sil.get("Riverside Health") == (1, 1) and sil.get("Lakeside Health") == (0, 0)
          and sil.get("Ungrouped") == (0, 0),
          f"sensors not reporting: per-type limits, broken sensor on a machine in use counted: {sil}")
    src = q(cur, "SELECT count(DISTINCT source) FROM daily_source_metrics")[0][0]
    check(src == 2, "daily_source_metrics splits spotlight vs hadrian")
    idm = dict(q(cur, "SELECT metric, value FROM daily_identity_metrics WHERE site_label='Riverside Health'"))
    check(idm.get("factor:WEAK_PASSWORD") == 1 and idm.get("mfa_registered") == 1, f"identity metrics {idm}")
    check((idm.get("cat:stale_accounts"), idm.get("cat:weak_passwords"), idm.get("cat:access_accounts"),
           idm.get("catgroup:exposure")) == (1, 1, 1, 2),
          "identity categories count distinct enabled accounts (stale + weak on one account = 1 in the group)")
    check((idm.get("total:enabled_accounts"), idm.get("total:enabled_people")) == (2, 0),
          f"enabled-account totals: service + VPN account, neither a person's main account: {idm}")

    print("== reports")
    out = os.path.join(tmp, "out")
    since = (dt.date.today() - dt.timedelta(days=6)).isoformat()
    check(run("exec_report.py", "--since", since, "--until", dt.date.today().isoformat(),
              "--out", os.path.join(out, "exec-region.html")) == 0, "exec_report.py region exited 0")
    check(run("exec_report.py", "--since", since, "--until", dt.date.today().isoformat(),
              "--site", "Riverside Health", "--out", os.path.join(out, "exec-site.html")) == 0,
          "exec_report.py single site exited 0")
    html_ = open(os.path.join(out, "exec-region.html")).read() if os.path.exists(os.path.join(out, "exec-region.html")) else ""
    check(html_.count('class="page"') == 6, "exec report has 6 pages")
    check("Removable Media (USB)" in html_ and "Business data written" in html_ and "medical images" in html_,
          "exec report has the removable-media page")
    site_html = open(os.path.join(out, "exec-site.html")).read() if os.path.exists(os.path.join(out, "exec-site.html")) else ""
    check("jnurse" not in site_html and "j*****" in site_html and "Ward 3" in site_html,
          "USB people masked in the site report, department shown")
    check("Most business data written" not in html_ and "j*****" not in html_, "region report has no per-person USB table")
    import exec_report
    if exec_report.find_chromium():
        pdf = os.path.join(out, "exec-region.pdf")
        check(os.path.exists(pdf) and open(pdf, "rb").read(5) == b"%PDF-", "exec report also rendered to PDF")
    else:
        print("  SKIP PDF check -- no Chromium on this machine")
    check("What we're protecting" in html_ and "Operating systems" in html_, "exec report has the estate page")
    check("someone@riversidehealth.test" not in html_ and "Infostealer" not in html_,
          "low-severity infostealer risk left out of the external table, account never shown unmasked")
    check("Inbound + internal emails" in html_ and "Delivered to users" in html_, "exec report has the mail-flow funnel")
    check("Indicators of compromise or attack" in html_ and "Reused passwords" in html_ and "WATCHED" not in html_,
          "exec report shows the identity categories, not raw factors")
    check(run("action_pack.py", "--out-dir", out, "--since", since, "--until", dt.date.today().isoformat()) == 0,
          "action_pack.py exited 0")
    packs = sorted(os.listdir(out)) if os.path.isdir(out) else []
    check(sum(f.endswith(".xlsx") for f in packs) == 4, f"action pack per site + region: {packs}")
    region_xlsx = next((os.path.join(out, f) for f in packs if f.startswith("region-")), None)
    if region_xlsx:
        from openpyxl import load_workbook
        wb = load_workbook(region_xlsx)
        check({"USB file writes", "USB drives", "USB malware"} <= set(wb.sheetnames), f"USB tabs: {wb.sheetnames}")
        ws = wb["USB file writes"]
        hdr = [c.value for c in ws[1]]
        top = dict(zip(hdr, [c.value for c in ws[2]]))
        check(top.get("Person") == "Jo Nurse" and "Email / mailbox data" in (top.get("Flags") or "")
              and "Large volume" in top["Flags"] and "Patient list.xlsx" in (top.get("Largest business files") or ""),
              f"named person first, flagged, with file names: {top}")
        drives = {r[2]: r for r in wb["USB drives"].iter_rows(min_row=2, values_only=True)}
        check("Shared drive" in (drives.get("SER1") or ("",) * 14)[13], f"SER1 used by two people flagged shared: {drives}")
        check(wb["USB malware"].max_row == 2, "USB malware tab lists the detection")
        nr = [dict(zip([c.value for c in wb["Sensors not reporting"][1]], [c.value for c in r]))
              for r in wb["Sensors not reporting"].iter_rows(min_row=2)]
        check(len(nr) == 1 and nr[0]["Hostname"] == "RANDOM01" and nr[0]["Likely"].startswith("Sensor broken")
              and nr[0]["Machine last seen (AD / network)"] is not None,
              f"sensors-not-reporting tab: the broken sensor on a machine in use, with when it was last seen: {nr}")
        acc = [dict(zip([c.value for c in wb["All accounts"][1]], [c.value for c in r]))
               for r in wb["All accounts"].iter_rows(min_row=2)]
        check(acc and acc[0]["Display name"] == "Svc Backup" and acc[0]["Stale"] == "Yes"
              and acc[0]["Password never expires"] == "Yes"
              and any(a["Title"] == "Registered Nurse" and a["Department"] == "Ward 3" for a in acc)
              and all(a["Enabled"] != "No" for a in acc),
              f"All accounts tab: enabled accounts only, stale first, with title/department: {acc[:2]}")

    print("== scorecard")
    conn.commit()   # end this script's read transaction: impact.py's ensure_schema needs table locks
    check(run("impact.py", "--only", "metrics") == 0, "impact.py metrics exited 0")
    sc_out = os.path.join(tmp, "scorecard")
    check(run("scorecard.py", "--out-dir", sc_out, "--month", dt.date.today().strftime("%Y-%m"), "--no-pdf") == 0,
          "scorecard.py exited 0")
    cards = sorted(os.listdir(sc_out)) if os.path.isdir(sc_out) else []
    check(len(cards) == 4 and all("-scorecard-" in f and f.endswith(".html") for f in cards),
          f"scorecard per site + region: {cards}")
    reg = next((os.path.join(sc_out, f) for f in cards if f.startswith("region-")), None)
    reg_html = open(reg).read() if reg else ""
    check(reg_html.count('<section class="page') == 2 and 'class="page wide"' in reg_html
          and "By health service" in reg_html
          and "Servers patched in the last 35 days" in reg_html and "Riverside Health" in reg_html,
          "region scorecard: measures page + landscape by-service table")
    one = next((os.path.join(sc_out, f) for f in cards if f.startswith("Riverside")), None)
    one_html = open(one).read() if one else ""
    check(one_html.count('<section class="page') == 1 and "Lakeside Health" not in one_html,
          "a site's scorecard shows only that site")
    check("Too early" in one_html, "one day of snapshots: directions are 'too early', not invented")
    import publish
    os.makedirs(os.path.join(tmp, "monthly-2026-11"), exist_ok=True)
    for f in ("Riverside_Health-scorecard-2026-11.pdf", "Riverside_Health-scorecard-2026-11.html",
              "Riverside_Health-actions-2026-11.xlsx"):
        open(os.path.join(tmp, "monthly-2026-11", f), "w").close()
    dest = sorted(d for _, d in publish.candidates(tmp))
    check(dest == ["Riverside_Health/Monthly/2026-11 Riverside_Health Action Pack.xlsx",
                   "Riverside_Health/Monthly/2026-11 Riverside_Health Security Scorecard.pdf"],
          f"publish files the scorecard PDF (not the HTML) next to the action pack: {dest}")

    print("== dashboard SQL")
    n = 0
    for path in sorted(glob.glob(os.path.join(ROOT, "grafana", "*.json"))):
        d = json.load(open(path))

        def walk(ps):
            for p in ps:
                yield p
                yield from walk(p.get("panels", []))
        site_vals = "'Riverside Health','Lakeside Health','Ungrouped'"
        for p in walk(d.get("panels", [])):
            for t in p.get("targets", []):
                sql = t.get("rawSql")
                if not sql:
                    continue
                sql = (sql.replace("$site", site_vals).replace("${site:sqlstring}", site_vals)
                       .replace("$source", "'falcon_spotlight','hadrian'")
                       .replace("$__timeFilter(", "TRUE OR (").replace("$__timeFrom()", "now()-interval '30 days'")
                       .replace("$__timeTo()", "now()")
                       .replace("${__from:date}", (dt.date.today() - dt.timedelta(days=30)).isoformat())
                       .replace("${__to:date}", dt.date.today().isoformat()))
                sql = re.sub(r"\$__timeGroupAlias\(([^,]+),[^)]+\)", r"date_trunc('day', \1) AS time", sql)
                sql = re.sub(r"\$__timeGroup\(([^,]+),[^)]+\)", r"date_trunc('day', \1)", sql)
                try:
                    cur.execute(sql)
                    cur.fetchall()
                    n += 1
                except Exception as e:
                    conn.rollback()
                    check(False, f"{os.path.basename(path)} :: {p.get('title')} :: {e}".splitlines()[0])
    check(n > 0, f"{n} dashboard queries executed cleanly")

    print(f"\n{'ALL PASSED' if not failures else f'{len(failures)} FAILURE(S)'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
