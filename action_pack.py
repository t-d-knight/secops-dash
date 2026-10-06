#!/usr/bin/env python3
"""
Monthly per-site action pack: the lists a service has to work through, as
one Excel workbook per site (plus a whole-of-region one with a Site column).
The exec report (exec_report.py) says how things are trending; this says
which accounts, hosts and external risks to fix.

Tabs:
  Summary              row counts per tab, as-at date, how to use it
  Stale accounts       Falcon Identity STALE_ACCOUNT
  Weak-compromised pw  WEAK_PASSWORD / CREDENTIAL_THEFT (same pair as the exec report tile)
  Duplicate passwords  DUPLICATE_PASSWORD
  No MFA               ACCOUNT_WITHOUT_MFA_CONFIGURED
  Stale sensors        managed hosts silent longer than reporting.stale_sensor_days
  Missing sensors      machines with no Falcon sensor: in-use AD computer accounts (servers,
                       DCs, workstations) first, then devices only seen on the network
  External risks       Hadrian: confirmed risks (verified / unpatched tech / infected
                       device) plus anything critical/high, incl. infostealer infections

Endpoint vulnerability tabs (top remediations, worst endpoints/servers) are
deliberately not here yet: Falcon vulnerability assessment currently covers
~9% of active hosts (2026-10-06), so those lists would name dormant machines.
Add them once coverage is fixed.

Account names are NOT masked here (unlike the exec report) -- the service
needs them to act. These files hold real usernames and compromised-account
details: distribute per site, to that site's contacts only.

    python3 action_pack.py --config config.yaml              # every site + region
    python3 action_pack.py --config config.yaml --site SITE-A

Writes reports/monthly-YYYY-MM/<site>-actions-YYYY-MM.xlsx (gitignored).
"""
import argparse
import datetime as dt
import os
import re
from typing import Any, Dict, List, Optional, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config as config_mod
import discovery_classes as dc
from db import pg_connect

SITE = "(%(site)s::text IS NULL OR {col} = %(site)s::text)"
HEADER_FILL = PatternFill("solid", fgColor="0F172A")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)

IDENTITY_COLS = ["Display name", "UPN", "Account (SAM)", "Domain", "OU", "Enabled", "Identity risk",
                 "Password last changed", "Created", "Other risk factors"]


def identity_sql(factors: Sequence[str], with_issue: bool = False) -> str:
    """Accounts carrying any of `factors`, enabled first, then by identity
    risk. "Other risk factors" lists everything else flagged on the account,
    so a service can see e.g. a stale account that also has no MFA."""
    flist = ", ".join(f"'{f}'" for f in factors)
    issue = (f"string_agg(DISTINCT f.factor_type, ', ') FILTER (WHERE f.factor_type IN ({flist})) AS issue, "
             if with_issue else "")
    return f"""
        SELECT e.site_label, {issue}e.display_name, e.upn, e.sam_account_name, e.domain, e.ou,
               CASE WHEN e.enabled THEN 'Yes' WHEN e.enabled IS FALSE THEN 'No' ELSE '?' END,
               initcap(e.risk_severity), e.password_last_change::date, e.account_created::date,
               (SELECT string_agg(o.factor_type, ', ' ORDER BY o.factor_type) FROM identity_risk_factors o
                 WHERE o.source = e.source AND o.entity_id = e.entity_id AND o.factor_type NOT IN ({flist}))
        FROM identity_entities e
        JOIN identity_risk_factors f ON f.source = e.source AND f.entity_id = e.entity_id
        WHERE NOT e.retired AND f.factor_type IN ({flist}) AND {SITE.format(col="e.site_label")}
        GROUP BY e.source, e.entity_id
        ORDER BY e.enabled IS NOT TRUE,
                 CASE e.risk_severity WHEN 'high' THEN 1 WHEN 'medium' THEN 2 WHEN 'low' THEN 3 ELSE 4 END,
                 e.display_name"""


TABS: List[Tuple[str, str, List[str], str]] = [
    # (tab name, sql, columns after Site, what to do)
    ("Stale accounts", identity_sql(["STALE_ACCOUNT"]), IDENTITY_COLS,
     "Accounts with no recent activity. Disable (or confirm still needed) -- enabled ones first."),
    ("Weak-compromised pw", identity_sql(["WEAK_PASSWORD", "CREDENTIAL_THEFT"], with_issue=True),
     ["Issue"] + IDENTITY_COLS,
     "Weak or stolen passwords. Force a reset; for CREDENTIAL_THEFT also review recent sign-ins."),
    ("Duplicate passwords", identity_sql(["DUPLICATE_PASSWORD"]), IDENTITY_COLS,
     "Accounts sharing a password with another account. Reset to unique passwords."),
    ("No MFA", identity_sql(["ACCOUNT_WITHOUT_MFA_CONFIGURED"]), IDENTITY_COLS,
     "Accounts with no MFA method registered. Enrol, or disable if unused."),
    ("Stale sensors", f"""
        SELECT site_label, hostname, product_type, os_version, last_seen::date,
               (now()::date - last_seen::date) AS days_silent, sensor_version,
               array_to_string(ous, ' / '), array_to_string(groups, ', ')
        FROM assets
        WHERE source = 'falcon' AND NOT retired AND last_seen < now() - (%(stale_days)s || ' days')::interval
          AND {SITE.format(col="site_label")}
        ORDER BY last_seen""",
     ["Hostname", "Type", "OS", "Last seen", "Days silent", "Sensor version", "OU", "Host groups"],
     "Managed hosts whose Falcon sensor hasn't checked in. Decommissioned: remove from Falcon. "
     "Still in use: fix the sensor -- these hosts are unprotected and unassessed."),
    ("Missing sensors", f"""
        SELECT site_label, {dc.sql_case()}, COALESCE(hostname, '(no hostname)'), os_version,
               array_to_string(ips, ', '), mac_vendor, last_seen::date, description, {dc.sql_case(part=1)}
        FROM assets
        WHERE source = 'falcon_unmanaged' AND NOT retired AND discovery_class IN {dc.sql_in(dc.ACTIONABLE)}
          AND {SITE.format(col="site_label")}
        ORDER BY CASE discovery_class WHEN 'domain_controller_no_sensor' THEN 1 WHEN 'server_no_sensor' THEN 2
                                      WHEN 'cloud_no_sensor' THEN 3 WHEN 'workstation_no_sensor' THEN 4
                                      WHEN 'vmware_nic' THEN 5 ELSE 6 END, hostname""",
     ["Class", "Hostname", "OS", "IPs", "MAC vendor", "Last seen", "AD description", "What to do"],
     "Machines with no Falcon sensor. Named servers/workstations are AD accounts in active use -- install a "
     "sensor or record why not. Network-only devices: identify by IP (PC, VM, or printer/IoT that can't take one)."),
    ("External risks", f"""
        SELECT vf.site_label, vf.vendor_priority, vf.risk_type, vf.plugin_family, vf.title, vf.hostname,
               vf.first_found::date, vf.state, left(vf.solution, 500)
        FROM vuln_findings vf
        WHERE vf.source = 'hadrian' AND vf.state IN ('OPEN','REOPENED')
          AND (vf.risk_type IN ('Verified','UnpatchedTechnology','InfectedDevice')
               OR vf.vendor_priority IN ('Critical','High'))
          AND {SITE.format(col="vf.site_label")}
        ORDER BY CASE WHEN vf.vendor_priority IN ('Critical','High') THEN 1
                      WHEN vf.risk_type = 'InfectedDevice' THEN 2 ELSE 3 END,
                 CASE vf.vendor_priority WHEN 'Critical' THEN 1 WHEN 'High' THEN 2 WHEN 'Medium' THEN 3
                                         WHEN 'Low' THEN 4 ELSE 5 END,
                 vf.first_found""",
     ["Severity", "Type", "Category", "Risk", "Asset", "First seen", "State", "Remediation"],
     "Internet-facing risks found by Hadrian: confirmed ones plus anything critical/high. "
     "Infostealer infections name the compromised accounts -- reset them and check the device."),
]


def safe_name(s: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", s)


def write_sheet(wb: Workbook, title: str, cols: List[str], rows: List[Tuple]) -> None:
    ws = wb.create_sheet(title[:31])
    ws.append(cols)
    for c in ws[1]:
        c.font, c.fill = HEADER_FONT, HEADER_FILL
    widths = [len(c) for c in cols]
    for r in rows:
        ws.append(list(r))
        for i, v in enumerate(r):
            widths[i] = max(widths[i], min(len(str(v)) if v is not None else 0, 60))
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w + 2
        if any(isinstance(r[i - 1], dt.date) for r in rows[:50]):
            for cell in ws[get_column_letter(i)][1:]:
                cell.number_format = "yyyy-mm-dd"
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions


def build_workbook(cur, site: Optional[str], as_at: dt.date, stale_days: int) -> Tuple[Workbook, Dict[str, int]]:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    counts: Dict[str, int] = {}
    region = site is None
    for name, sql, cols, _ in TABS:
        cur.execute(sql, {"site": site, "stale_days": stale_days})
        rows = cur.fetchall()
        # region workbook keeps the Site column; a site's own workbook drops it
        rows = [r if region else r[1:] for r in rows]
        write_sheet(wb, name, (["Site"] if region else []) + cols, rows)
        counts[name] = len(rows)

    summary["A1"] = f"Security action pack -- {'Whole of region' if region else site}"
    summary["A1"].font = TITLE_FONT
    summary["A2"] = f"As at {as_at:%d %B %Y}. Point-in-time lists from the secops dashboard's latest collection."
    summary.append([])
    summary.append(["List", "Items", "What to do"])
    for c in summary[4]:
        c.font, c.fill = HEADER_FONT, HEADER_FILL
    for name, _, _, todo in TABS:
        summary.append([name, counts[name], todo])
    summary.append([])
    for line in (
        "Endpoint vulnerability lists (top remediations, worst endpoints/servers) will be added once "
        "Falcon vulnerability assessment covers the whole estate.",
        "This workbook contains real account names and compromised-credential details. "
        "Keep it within the site's security contacts.",
    ):
        summary.append([line])
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 8
    summary.column_dimensions["C"].width = 110
    for row in summary.iter_rows(min_row=5):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    return wb, counts


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--config", default="config.yaml")
    ap.add_argument("--site", help="one site (its label); default is every site plus the region workbook")
    ap.add_argument("--out-dir", help="default: reports/monthly-YYYY-MM")
    args = ap.parse_args()

    cfg = config_mod.load_config(args.config)
    labels = [s["label"] for s in cfg.get("sites", [])] + [cfg.get("ungrouped_label", "Ungrouped")]
    if args.site and args.site not in labels:
        ap.error(f"unknown site {args.site!r}; configured: {', '.join(labels)}")
    as_at = dt.date.today()
    stale_days = int(cfg.get("reporting", {}).get("stale_sensor_days", 7))
    out_dir = args.out_dir or os.path.join("reports", f"monthly-{as_at:%Y-%m}")
    os.makedirs(out_dir, exist_ok=True)

    conn = pg_connect(cfg)
    cur = conn.cursor()
    jobs: List[Optional[str]] = [args.site] if args.site else [None] + labels
    for site in jobs:
        wb, counts = build_workbook(cur, site, as_at, stale_days)
        path = os.path.join(out_dir, f"{safe_name(site or 'region')}-actions-{as_at:%Y-%m}.xlsx")
        wb.save(path)
        print(f"[action_pack] wrote {path} " + ", ".join(f"{k}={v}" for k, v in counts.items()))
    conn.close()


if __name__ == "__main__":
    main()
