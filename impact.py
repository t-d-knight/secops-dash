#!/usr/bin/env python3
"""
Programme impact: what changed since this started, for the Impact dashboard
and the monthly Impact report.

  daily_impact_metrics  a curated set of counts per site, every night, computed
                        with the exec report's own fetch functions so the
                        numbers match the weekly reports. The first day on
                        record is "Day 0" (impact.baseline_date overrides).
  patch_cohort_metrics  each month's Patch Tuesday Microsoft findings (first
                        found in the week after release, on hosts already being
                        assessed) and how many closed within 7 / 14 / 30 / 60
                        days -- time to remediate without needing history, and
                        comparable month on month.
  impact_actions        the action log: what was found, how, and what was done.
                        Kept as an Excel workbook on the publish share
                        (impact.action_log, created with headers if missing) so
                        it can be edited in SharePoint / Excel; loaded nightly.

Rates are always derived from counts at display time (SCORECARD), so a
region figure is the sum of sites, never an average of percentages.

    python3 impact.py --config config.yaml              # today's metrics, cohorts, action log
    python3 impact.py --config config.yaml --only actions
"""
import argparse
import datetime as dt
import io
import sys
from typing import Any, Dict, List, Optional, Tuple

import config as config_mod
import db_schema
import exec_report as er
import os_lifecycle
from db import pg_connect

from impact_spec import METRICS, SCORECARD  # noqa: F401 (re-exported)

SRV_ACTIVE = ("FROM assets a LEFT JOIN host_patch_status p ON p.aid = a.source_asset_id "
              "LEFT JOIN host_insights i ON i.aid = a.source_asset_id "
              "WHERE a.source = 'falcon' AND NOT a.retired AND a.product_type IN ('server','domain_controller') "
              "AND a.last_seen > now() - interval '7 days' AND a.site_label = %(site)s")


def site_metrics(cur, cfg: Dict[str, Any], site: str, p: er.Period, days_last_seen: int) -> Dict[str, int]:
    es = er.fetch_estate(cur, site)
    v = er.fetch_vulns(cur, p, site, days_last_seen)
    idn = er.fetch_identity(cur, site)
    a = er.fetch_alerts(cur, p, site)
    x = er.fetch_external(cur, p, site)
    cur.execute(f"""
        SELECT count(*), count(*) FILTER (WHERE p.last_patch_confirmed >= now() - interval '35 days'),
               count(*) FILTER (WHERE p.last_patch_confirmed < now() - interval '90 days'),
               count(*) FILTER (WHERE a.last_reboot < now() - interval '35 days'),
               count(*) FILTER (WHERE a.product_type = 'domain_controller'),
               count(*) FILTER (WHERE a.product_type = 'domain_controller' AND (i.os_is_eol
                   OR p.last_patch_confirmed IS NULL OR p.last_patch_confirmed < now() - interval '35 days'))
        {SRV_ACTIVE}""", {"site": site})
    srv = cur.fetchone()
    cur.execute("""
        SELECT count(*) FILTER (WHERE r.open_kev_count > 0)
        FROM assets a LEFT JOIN dash_asset_risk r ON r.source = 'falcon_spotlight' AND r.source_asset_id = a.source_asset_id
        WHERE a.source = 'falcon' AND NOT a.retired AND a.last_seen > now() - interval '7 days' AND a.site_label = %(site)s""",
                {"site": site})
    kev_hosts = cur.fetchone()[0]
    return {
        "servers_active": srv[0], "servers_patched_35d": srv[1], "servers_unpatched_90d": srv[2],
        "servers_up_35d": srv[3], "dcs_active": srv[4], "dcs_attention": srv[5],
        "hosts_active": es["active"], "hosts_with_kev": kev_hosts, "kev_open": v["kev_open"],
        "crit_open": v["open_crit_windowed"], "crit_past_sla": v["sla_breach_crit"],
        "high_open": v["open_high_windowed"], "high_past_sla": v["sla_breach_high"],
        "unsupported_os": es["unsupported"], "missing_sensors": es["missing_named"],
        "stale_sensors": v["stale_sensors"], "managed_hosts": v["managed_hosts"],
        "enabled_accounts": idn["totals"]["enabled_accounts"], "stale_accounts": idn["cats"]["stale_accounts"],
        "people_accounts": idn["totals"]["enabled_people"], "people_no_mfa": idn["cats"]["no_mfa"],
        "unused_access_accounts": idn["stale_access"], "compromise_indicators": idn["groups"]["compromise"],
        "identity_triage": a["open_identity_hm"],
        "ext_confirmed_open": x["confirmed"], "ext_crit_high_open": x["crit_high_any"], "leaked_creds": x["leaked_creds"],
    }


def rollup_metrics(conn, cfg: Dict[str, Any], snapshot_date: dt.date) -> int:
    rep = cfg.get("reporting") or {}
    er.OS_EOS.clear()
    er.OS_EOS.update(os_lifecycle.table(rep.get("os_end_of_support")))
    p = er.week_period(snapshot_date)
    cur = conn.cursor()
    rows = []
    for site in er.site_labels(cfg):
        for k, val in site_metrics(cur, cfg, site, p, rep.get("days_last_seen", 30)).items():
            rows.append((snapshot_date, site, k, int(val or 0)))
    cur.execute("DELETE FROM daily_impact_metrics WHERE snapshot_date = %s", (snapshot_date,))
    cur.executemany("INSERT INTO daily_impact_metrics (snapshot_date, site_label, metric, value) VALUES (%s, %s, %s, %s)",
                    rows)
    conn.commit()
    return len(rows)


# ------------------------------------------------------------------ cohorts
def patch_tuesday(year: int, month: int) -> dt.date:
    """Second Tuesday (US release date; lands here the next morning)."""
    d = dt.date(year, month, 8)
    return d + dt.timedelta(days=(1 - d.weekday()) % 7)


def rollup_cohorts(conn, cfg: Dict[str, Any], today: dt.date) -> int:
    """Recompute every cohort from impact.cohort_start (default 2026-10) to now."""
    imp = cfg.get("impact") or {}
    start = dt.date.fromisoformat(str(imp.get("cohort_start", "2026-10-01")))
    vendor = imp.get("cohort_vendor", "microsoft")
    window = int(imp.get("cohort_window_days", 7))
    cur = conn.cursor()
    # when each host was first assessed -- a cohort only counts hosts assessed before its release
    cur.execute("""CREATE TEMP TABLE IF NOT EXISTS cohort_host_first AS
                   SELECT source_asset_id, min(first_found) AS f FROM vuln_findings
                   WHERE source = 'falcon_spotlight' GROUP BY 1""")
    cur.execute("CREATE INDEX IF NOT EXISTS cohort_host_first_i ON cohort_host_first (source_asset_id)")
    n = 0
    y, m = start.year, start.month
    while True:
        release = patch_tuesday(y, m)
        lo = release + dt.timedelta(days=1)          # arrives here the morning after
        if lo > today:
            break
        hi = lo + dt.timedelta(days=window)
        cur.execute("""
            WITH c AS (
                SELECT vf.site_label, vf.source_asset_id, vf.state, vf.last_fixed
                FROM vuln_findings vf JOIN cohort_host_first h ON h.source_asset_id = vf.source_asset_id
                WHERE vf.source = 'falcon_spotlight' AND vf.product_vendor = %(vendor)s
                  AND vf.first_found >= %(lo)s AND vf.first_found < %(hi)s AND h.f < %(lo)s)
            SELECT site_label, count(*), count(DISTINCT source_asset_id),
                   count(*) FILTER (WHERE state = 'FIXED' AND last_fixed < %(lo)s::date + 7),
                   count(*) FILTER (WHERE state = 'FIXED' AND last_fixed < %(lo)s::date + 14),
                   count(*) FILTER (WHERE state = 'FIXED' AND last_fixed < %(lo)s::date + 30),
                   count(*) FILTER (WHERE state = 'FIXED' AND last_fixed < %(lo)s::date + 60),
                   count(*) FILTER (WHERE state = 'FIXED'), count(*) FILTER (WHERE state = 'EXPIRED')
            FROM c GROUP BY 1""", {"vendor": vendor, "lo": lo, "hi": hi})
        rows = cur.fetchall()
        cur.execute("DELETE FROM patch_cohort_metrics WHERE cohort = %s", (release,))
        elapsed = (today - lo).days
        cur.executemany("""
            INSERT INTO patch_cohort_metrics (cohort, site_label, findings, hosts, closed_7, closed_14, closed_30,
                                              closed_60, closed_now, expired, days_elapsed)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)""",
                        [(release, r[0], r[1], r[2],
                          r[3] if elapsed >= 7 else None, r[4] if elapsed >= 14 else None,
                          r[5] if elapsed >= 30 else None, r[6] if elapsed >= 60 else None,
                          r[7], r[8], elapsed) for r in rows])
        n += len(rows)
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    conn.commit()
    return n


# -------------------------------------------------------------- action log
ACTION_COLS = ["Found", "What was found", "How it was found", "Site", "Severity", "Action taken", "Owner team",
               "Status", "Closed", "Ticket ref", "Notes"]
SOURCES = ["Weekly report", "Quarterly report", "Servers report", "Action pack", "Dashboard", "Identity digest",
           "Alert", "Other"]
SEVERITIES = ["Critical", "High", "Medium", "Low"]
STATUSES = ["Open", "In progress", "Done", "Accepted risk"]


def action_log_template() -> bytes:
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation
    from action_pack import HEADER_FILL, HEADER_FONT
    wb = Workbook()
    ws = wb.active
    ws.title = "Actions"
    ws.append(ACTION_COLS)
    for c in ws[1]:
        c.font, c.fill = HEADER_FONT, HEADER_FILL
    for col, w in zip("ABCDEFGHIJK", (12, 60, 18, 12, 10, 50, 18, 13, 12, 14, 40)):
        ws.column_dimensions[col].width = w
    for col, opts in (("C", SOURCES), ("E", SEVERITIES), ("H", STATUSES)):
        dv = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}2000")
    ws.freeze_panes = "A2"
    help_ws = wb.create_sheet("How to use")
    for line in ("One row per thing this programme surfaced that led (or should lead) to action.",
                 "Found / Closed: dates. How it was found / Severity / Status: pick from the list.",
                 "Status Done or Accepted risk counts as resolved. Ticket ref: SDP request number, if any.",
                 "Loaded nightly into the Impact dashboard and the monthly Impact report. Don't rename the "
                 "'Actions' sheet or the column headings."):
        help_ws.append([line])
    help_ws.column_dimensions["A"].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell_date(v: Any) -> Optional[dt.date]:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str) and v.strip():
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d %b %Y", "%d-%m-%Y"):
            try:
                return dt.datetime.strptime(v.strip(), fmt).date()
            except ValueError:
                pass
    return None


def load_actions(conn, cfg: Dict[str, Any]) -> str:
    """Read the action log workbook from the share into impact_actions (replacing
    it). Creates the workbook with headers if it doesn't exist yet."""
    from openpyxl import load_workbook
    import publish
    rel = ((cfg.get("impact") or {}).get("action_log") or "Impact/Impact action log.xlsx").strip("/")
    t = publish.target(cfg.get("publish") or {})
    if not hasattr(t, "smb"):
        return "action log: publish.method isn't smb -- skipped"
    path = t._unc(rel)
    if not t.smb.path.exists(path):
        t.smb.makedirs(path.rsplit("\\", 1)[0], exist_ok=True)
        with t.smb.open_file(path, mode="wb") as fh:
            fh.write(action_log_template())
        return f"action log: created {path}"
    with t.smb.open_file(path, mode="rb") as fh:
        wb = load_workbook(io.BytesIO(fh.read()), data_only=True)
    ws = wb["Actions"] if "Actions" in wb.sheetnames else wb.active
    head = [str(c.value or "").strip() for c in ws[1]]
    idx = {h: head.index(h) for h in ACTION_COLS if h in head}
    if "What was found" not in idx:
        return f"action log: no 'What was found' column in {path} -- left as it was"
    rows = []
    for n, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        get = lambda h: r[idx[h]] if h in idx and idx[h] < len(r) else None
        what = get("What was found")
        if not what or not str(what).strip():
            continue
        status = str(get("Status") or "Open").strip()
        rows.append((n, _cell_date(get("Found")), str(what).strip(), get("How it was found"), get("Site"),
                     get("Severity"), get("Action taken"), get("Owner team"), status,
                     _cell_date(get("Closed")), get("Ticket ref"), get("Notes")))
    cur = conn.cursor()
    cur.execute("DELETE FROM impact_actions")
    cur.executemany("""
        INSERT INTO impact_actions (row_no, found, what, source, site_label, severity, action, owner_team, status,
                                    closed, ticket_ref, notes)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)""",
                    [tuple(None if v is None else (v if not isinstance(v, str) else v.strip()) for v in row)
                     for row in rows])
    conn.commit()
    return f"action log: {len(rows)} rows from {path}"


# --------------------------------------------------------------------- main
def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--config", default="config.yaml")
    ap.add_argument("--only", choices=["metrics", "cohorts", "actions"])
    args = ap.parse_args()
    cfg = config_mod.load_config(args.config)
    db_schema.ensure_schema(cfg)
    conn = pg_connect(cfg)
    today = dt.date.today()
    rc = 0
    try:
        if args.only in (None, "metrics"):
            print(f"[impact] metrics: {rollup_metrics(conn, cfg, today)} rows for {today}")
        if args.only in (None, "cohorts"):
            print(f"[impact] cohorts: {rollup_cohorts(conn, cfg, today)} rows")
        if args.only in (None, "actions"):
            try:
                print(f"[impact] {load_actions(conn, cfg)}")
            except Exception as e:      # the share being down mustn't fail the nightly run
                conn.rollback()
                print(f"[impact] action log not loaded: {type(e).__name__}: {e}", file=sys.stderr)
                rc = 1
    finally:
        conn.close()
    return rc


if __name__ == "__main__":
    sys.exit(main())
